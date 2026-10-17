from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Project directories
BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "outputs"
EXCEL_DIR = BASE_DIR / "06_Excel"

EXCEL_DIR.mkdir(parents=True, exist_ok=True)

# Input CSV files
reports = {
    "Portfolio Analysis": "bond_portfolio_analysis.csv",
    "Price Predictions": "bond_price_predictions.csv",
    "Sensitivity Analysis": "bond_sensitivity_analysis.csv",
    "Risk Summary": "portfolio_risk_summary.csv",
}

# Create Excel workbook
excel_path = EXCEL_DIR / "Bond_Portfolio_Analysis.xlsx"

with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
    for sheet_name, filename in reports.items():
        csv_path = OUTPUT_DIR / filename

        if not csv_path.exists():
            print(f"WARNING: Missing file: {csv_path}")
            continue

        df = pd.read_csv(csv_path)
        df.to_excel(writer, sheet_name=sheet_name, index=False)

# Format workbook
workbook = load_workbook(excel_path)

for worksheet in workbook.worksheets:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        cell.alignment = Alignment(horizontal="center")

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2, 35
        )

workbook.save(excel_path)

print("\nExcel report created successfully!")
print(f"Location: {excel_path}")
print(f"Worksheets: {workbook.sheetnames}")